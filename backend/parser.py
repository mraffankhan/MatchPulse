"""
Resume and Document Parser.
Handles PDF, DOCX, and TXT files, extracting raw text, structured sections, contact details,
work experience duration, seniority, and education level.
"""

import io
import re
from typing import Dict, Any, List, Optional
import pypdf
import docx

from .taxonomy import (
    ALL_SKILL_LOOKUPS,
    canonicalize_skill,
    get_skill_category,
    EDUCATION_LEVELS,
    SENIORITY_RANKS
)


class ResumeParser:
    """Robust extractor for multi-format resumes and job descriptions."""

    SECTION_PATTERNS = {
        "summary": [
            r"\b(professional\s+summary|summary|profile|about\s+me|objective|overview)\b"
        ],
        "skills": [
            r"\b(technical\s+skills|skills\s*(&|and)?\s*competencies|skills|technologies|core\s+competencies|toolset|proficiencies)\b"
        ],
        "experience": [
            r"\b(work\s+experience|professional\s+experience|experience|employment\s+history|career\s+history|work\s+history)\b"
        ],
        "education": [
            r"\b(education|academic\s+background|academic\s+qualifications|academics|qualifications)\b"
        ],
        "projects": [
            r"\b(projects|personal\s+projects|key\s+projects|portfolio\s+projects)\b"
        ],
        "certifications": [
            r"\b(certifications|licenses\s*(&|and)?\s*certifications|certificates|accreditations)\b"
        ]
    }

    @staticmethod
    def extract_text_from_pdf(file_bytes: bytes) -> str:
        """Extract text content from PDF bytes using pypdf."""
        try:
            reader = pypdf.PdfReader(io.BytesIO(file_bytes))
            text_parts = []
            for page in reader.pages:
                page_text = page.extract_text()
                if page_text:
                    text_parts.append(page_text)
            return "\n".join(text_parts).strip()
        except Exception as e:
            raise ValueError(f"Failed to parse PDF: {str(e)}")

    @staticmethod
    def extract_text_from_docx(file_bytes: bytes) -> str:
        """Extract text content from DOCX bytes using python-docx."""
        try:
            doc = docx.Document(io.BytesIO(file_bytes))
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs).strip()
        except Exception as e:
            raise ValueError(f"Failed to parse DOCX: {str(e)}")

    @classmethod
    def parse_document(cls, filename: str, content_bytes: bytes) -> str:
        """Auto-detect format and extract text."""
        fname = filename.lower()
        if fname.endswith(".pdf"):
            return cls.extract_text_from_pdf(content_bytes)
        elif fname.endswith(".docx"):
            return cls.extract_text_from_docx(content_bytes)
        elif fname.endswith(".txt") or fname.endswith(".md"):
            return content_bytes.decode("utf-8", errors="ignore")
        else:
            # Fallback to UTF-8 decode
            return content_bytes.decode("utf-8", errors="ignore")

    @classmethod
    def extract_contact_info(cls, text: str) -> Dict[str, Any]:
        """Extract email, phone, and professional profile links."""
        # Email
        email_pattern = r"[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+"
        emails = re.findall(email_pattern, text)
        email = emails[0] if emails else None

        # Phone (international & standard US/India/EU formats)
        phone_pattern = r"(?:\+?\d{1,3}[-.\s]?)?\(?\d{2,4}\)?[-.\s]?\d{3,4}[-.\s]?\d{3,4}"
        phones = re.findall(phone_pattern, text)
        clean_phones = [p.strip() for p in phones if len(re.sub(r"\D", "", p)) >= 10]
        phone = clean_phones[0] if clean_phones else None

        # Links (LinkedIn, GitHub)
        linkedin_match = re.search(r"(?:https?://)?(?:www\.)?linkedin\.com/in/([a-zA-Z0-9_-]+)", text, re.IGNORECASE)
        github_match = re.search(r"(?:https?://)?(?:www\.)?github\.com/([a-zA-Z0-9_-]+)", text, re.IGNORECASE)

        return {
            "email": email,
            "phone": phone,
            "linkedin": linkedin_match.group(0) if linkedin_match else None,
            "github": github_match.group(0) if github_match else None
        }

    @classmethod
    def extract_name(cls, text: str) -> Optional[str]:
        """Infer candidate name from the top lines of the document."""
        lines = [line.strip() for line in text.split("\n") if line.strip()]
        for line in lines[:5]:
            # Skip if contains email, links, or section headers
            if "@" in line or "http" in line or "github" in line or "linkedin" in line or "resume" in line.lower():
                continue
            # Words count check (typically 2-4 words)
            words = line.split()
            if 1 < len(words) <= 4 and all(w.isalpha() or "." in w for w in words):
                return line.title()
        return "Candidate"

    @classmethod
    def segment_sections(cls, text: str) -> Dict[str, str]:
        """Split document into canonical sections based on heading detection."""
        lines = text.split("\n")
        sections: Dict[str, List[str]] = {
            "header": [],
            "summary": [],
            "skills": [],
            "experience": [],
            "education": [],
            "projects": [],
            "certifications": [],
            "other": []
        }

        current_section = "header"

        for line in lines:
            line_str = line.strip()
            if not line_str:
                continue

            # Detect section header
            matched_section = None
            if len(line_str) < 50:
                for sec_key, patterns in cls.SECTION_PATTERNS.items():
                    for pattern in patterns:
                        if re.search(f"^{pattern}\\s*:?$", line_str, re.IGNORECASE) or \
                           re.search(f"^#+\\s*{pattern}", line_str, re.IGNORECASE):
                            matched_section = sec_key
                            break
                    if matched_section:
                        break

            if matched_section:
                current_section = matched_section
            else:
                sections[current_section].append(line_str)

        return {k: "\n".join(v).strip() for k, v in sections.items()}

    @classmethod
    def extract_skills(cls, text: str) -> List[Dict[str, str]]:
        """
        Extract all known technical & soft skills from text with canonical naming,
        deduplication, and domain categorization.
        """
        text_lower = f" {text.lower()} "
        # Replace common punctuation with spaces to enable clean boundary checks,
        # but preserve '.', '#', '+', '/' for C++, C#, .NET, CI/CD, React.js, etc.
        sanitized = re.sub(r"[,;:\(\)\[\]\{\}\*•|!\?]", " ", text_lower)
        sanitized = re.sub(r"\s+", " ", sanitized)

        found_skills = {}

        for term, _ in ALL_SKILL_LOOKUPS:
            # Word boundary regex ensuring we don't accidentally match substrings
            # Example: "c" won't match "cat", but "c++" will match "c++"
            escaped_term = re.escape(term)
            pattern = rf"(?<![a-zA-Z0-9]){escaped_term}(?![a-zA-Z0-9])"

            if re.search(pattern, sanitized):
                canonical = canonicalize_skill(term)
                if canonical not in found_skills:
                    found_skills[canonical] = {
                        "name": canonical,
                        "category": get_skill_category(canonical),
                        "matched_term": term
                    }

        return list(found_skills.values())

    @classmethod
    def extract_experience_years(cls, text: str) -> float:
        """
        Estimate total years of experience from explicit mentions
        (e.g., '5+ years of experience', 'Over 7 years') or date range spans.
        """
        explicit_patterns = [
            r"(\d+(?:\.\d+)?)\+?\s*(?:years|yrs)\b(?:\s*of\s*(?:relevant\s*)?experience)?",
            r"(?:over|more\s+than)\s*(\d+(?:\.\d+)?)\s*(?:years|yrs)"
        ]

        found_years = []
        for pat in explicit_patterns:
            matches = re.findall(pat, text, re.IGNORECASE)
            for m in matches:
                try:
                    val = float(m)
                    if 0.5 <= val <= 35:  # realistic bounds
                        found_years.append(val)
                except ValueError:
                    continue

        if found_years:
            return max(found_years)

        # Fallback: Count year ranges like (2018 - 2023, 2021 - Present)
        year_range_pattern = r"\b(20[0-2][0-9]|199[0-9])\s*[-–—to]+\s*(20[0-2][0-9]|present|current|now)\b"
        ranges = re.findall(year_range_pattern, text, re.IGNORECASE)
        total_span = 0.0
        current_year = 2026

        for start, end in ranges:
            try:
                start_yr = int(start)
                end_yr = current_year if end.lower() in ["present", "current", "now"] else int(end)
                if 0 <= (end_yr - start_yr) <= 25:
                    total_span += (end_yr - start_yr)
            except ValueError:
                pass

        return min(max(total_span, 0.0), 30.0)

    @classmethod
    def extract_seniority(cls, text: str, given_title: Optional[str] = None) -> str:
        """
        Detect seniority level, prioritizing the title/headline and header lines
        before falling back to full-text scan.
        """
        # 1. First inspect the explicit title or top lines
        header_text = (given_title or "") + "\n" + "\n".join(text.split("\n")[:4])
        header_lower = header_text.lower()

        # Check for explicit seniority in header
        for title, rank in sorted(SENIORITY_RANKS.items(), key=lambda x: x[1], reverse=True):
            if title == "architect":
                # Ensure it's used as a role title, not just verb
                if re.search(r"\b(software|cloud|solution|enterprise|systems?|lead)\s+architect\b|\barchitect\s+(level|role|engineer)\b", header_lower):
                    return "Architect"
            elif title == "mid":
                if re.search(r"\b(mid-level|mid level|intermediate)\b", header_lower):
                    return "Mid-Level"
            else:
                pattern = rf"\b{re.escape(title)}\b"
                if re.search(pattern, header_lower):
                    return title.title()

        # 2. If not found in headers, scan full text with strict patterns
        text_lower = text.lower()
        detected_rank = 0
        detected_level = "Mid-Level"

        for title, rank in SENIORITY_RANKS.items():
            if title == "architect":
                pattern = r"\b(software|cloud|solution|enterprise|systems?)\s+architect\b"
            elif title == "mid":
                pattern = r"\b(mid-level|mid level|intermediate)\b"
            else:
                pattern = rf"\b{re.escape(title)}\s+(?:engineer|developer|architect|designer|scientist|lead|manager|specialist|consultant)\b|\b{re.escape(title)}\b"

            if re.search(pattern, text_lower):
                if rank > detected_rank:
                    detected_rank = rank
                    detected_level = title.title()

        return detected_level

    @classmethod
    def extract_education(cls, text: str) -> str:
        """Detect the highest education level mentioned."""
        text_lower = text.lower()
        detected_rank = 0
        detected_level = "Not Specified"

        for degree, rank in EDUCATION_LEVELS.items():
            pattern = rf"\b{re.escape(degree)}\b"
            if re.search(pattern, text_lower):
                if rank > detected_rank:
                    detected_rank = rank
                    detected_level = degree.upper() if len(degree) <= 4 else degree.title()

        return detected_level

    @classmethod
    def analyze_resume(cls, text: str, filename: str = "resume.txt") -> Dict[str, Any]:
        """Full structural and semantic parse of a candidate resume."""
        contact = cls.extract_contact_info(text)
        name = cls.extract_name(text)
        sections = cls.segment_sections(text)
        skills = cls.extract_skills(text)
        years = cls.extract_experience_years(text)
        seniority = cls.extract_seniority(text)
        education = cls.extract_education(text)

        word_count = len(text.split())

        return {
            "filename": filename,
            "candidate_name": name,
            "contact": contact,
            "sections": sections,
            "skills": skills,
            "skill_names": [s["name"] for s in skills],
            "experience_years": round(years, 1),
            "seniority_level": seniority,
            "education_level": education,
            "word_count": word_count,
            "raw_text": text
        }

    @classmethod
    def analyze_job(cls, text: str, title: str = "Target Job") -> Dict[str, Any]:
        """Full parse of a job description."""
        skills = cls.extract_skills(text)
        years = cls.extract_experience_years(text)
        seniority = cls.extract_seniority(text, given_title=title)
        education = cls.extract_education(text)

        return {
            "title": title,
            "skills": skills,
            "skill_names": [s["name"] for s in skills],
            "required_years": round(years, 1),
            "required_seniority": seniority,
            "required_education": education,
            "raw_text": text
        }
